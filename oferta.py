"""
Geração do Excel da Oferta Hospitalar por ULS

Folha 1 — Oferta por ULS:
  Capacidade Real = Camas Totais × Taxa de Ocupação
  (nº médio de camas efetivamente ocupadas/em uso, por dia, em 2025)

Fontes (todas referentes a 2025):
- lotacao-praticada-por-tipo-de-cama.xlsx  → nº de camas por instituição e tipo (mensal)
- taxa_de_ocupação.xlsx                    → taxa de ocupação por instituição (mensal, 0-1)
"""

import pandas as pd
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

LOTACAO_FILE = '/Users/margaridaclaro/Desktop/Oferta/lotacao-praticada-por-tipo-de-cama.xlsx'
OCUP_FILE    = '/Users/margaridaclaro/Desktop/oferta/taxa_de_ocupação.xlsx'
OUTPUT_FILE  = '/Users/margaridaclaro/Desktop/Oferta/Excel_Oferta_Portugal.xlsx'

C_HEADER = "1F3864"
C_SUBHDR = "2E75B6"
C_ALT    = "D9E1F2"
C_WHITE  = "FFFFFF"

ANO_ANALISE = 2025


def norm_name(nome):
    """Normaliza nome de instituição para matching: remove EPE/PPP, pontuação, espaços duplos, lowercase."""
    nome = str(nome)
    nome = re.sub(r',?\s*E\.?P\.?E\.?$', '', nome, flags=re.IGNORECASE)
    nome = re.sub(r',?\s*PPP$',          '', nome, flags=re.IGNORECASE)
    nome = nome.replace(' / ', ' ').replace('/', ' ').replace('-', ' ')
    return re.sub(r'\s+', ' ', nome).strip().lower()


# Correção de nomes de ULS divergentes entre o ficheiro de lotação e o de taxa de ocupação
# Chave = Key_Norm como sai da lotação; Valor = Key_Norm canónico (como sai da ocupação)
ULS_LOTACAO_TO_OCUP = {
    'unidade local de saúde castelo branco':              'unidade local de saúde de castelo branco',
    'unidade local de saúde de dão lafões':                'unidade local de saúde de viseu dão lafões',
    'unidade local de saúde de vila nova de gaia espinho': 'unidade local de saúde de gaia espinho',
}

def fix_key(key_norm):
    return ULS_LOTACAO_TO_OCUP.get(key_norm, key_norm)


# ══════════════════════════════════════════════════════════════════════════
# 1. CAMAS TOTAIS POR ULS (2025) — soma dos 4 tipos, média mensal anualizada
# ══════════════════════════════════════════════════════════════════════════

def load_camas():
    df = pd.read_excel(LOTACAO_FILE, sheet_name='Feuil1', header=0)
    df = df[['Período', 'Instituição', 'Tipo de Camas', 'Lotação']].copy()
    df['Ano'] = df['Período'].astype(str).str[:4].astype(int)
    df = df[df['Ano'] == ANO_ANALISE]
    df['Lotação'] = pd.to_numeric(df['Lotação'], errors='coerce').fillna(0)

    # Filtrar só Unidades Locais de Saúde
    df = df[df['Instituição'].str.contains('Unidade Local de Saúde', na=False)].copy()
    df['Key_Norm'] = df['Instituição'].apply(norm_name).apply(fix_key)
    por_mes = df.groupby(['Key_Norm', 'Instituição', 'Período'])['Lotação'].sum().reset_index()

    # Média mensal de camas totais ao longo de 2025 (mais robusto que usar só um mês)
    camas = por_mes.groupby(['Key_Norm', 'Instituição'])['Lotação'].mean().reset_index()
    camas.columns = ['Key_Norm', 'Instituicao_Lotacao', 'Camas_Totais_Media']
    camas['Camas_Totais_Media'] = camas['Camas_Totais_Media'].round(1)

    return camas[['Key_Norm', 'Instituicao_Lotacao', 'Camas_Totais_Media']]


# ══════════════════════════════════════════════════════════════════════════
# 2. TAXA DE OCUPAÇÃO POR ULS (2025) — média anual
# ══════════════════════════════════════════════════════════════════════════

def load_taxa_ocupacao():
    df = pd.read_excel(OCUP_FILE, sheet_name='folha1', header=7)
    df = df[df['Ano'] == ANO_ANALISE].copy()
    df['Valor'] = pd.to_numeric(df['Valor'], errors='coerce')
    df = df[df['Instituição'].str.contains('Unidade Local de Saúde', na=False)].copy()
    df['Key_Norm'] = df['Instituição'].apply(norm_name).apply(fix_key)

    agg = df.groupby('Key_Norm')['Valor'].mean().reset_index()
    agg.columns = ['Key_Norm', 'Taxa_Ocupacao_Media']
    agg['Taxa_Ocupacao_Media_Pct'] = (agg['Taxa_Ocupacao_Media'] * 100).round(1)

    return agg[['Key_Norm', 'Taxa_Ocupacao_Media', 'Taxa_Ocupacao_Media_Pct']]


# ══════════════════════════════════════════════════════════════════════════
# 3. CALCULAR OFERTA (CAPACIDADE REAL)
# ══════════════════════════════════════════════════════════════════════════

def build_oferta():
    camas = load_camas()
    ocupacao = load_taxa_ocupacao()

    merged = camas.merge(ocupacao, on='Key_Norm', how='left')

    # Capacidade Real (camas efetivamente ocupadas/dia, em média, em 2025)
    # Fórmula: Camas Totais × Taxa de Ocupação
    merged['Capacidade_Real'] = (
        merged['Camas_Totais_Media'] * merged['Taxa_Ocupacao_Media']
    ).round(1)

    # Diagnóstico de matching incompleto
    sem_ocup = merged[merged['Taxa_Ocupacao_Media'].isna()]['Instituicao_Lotacao'].tolist()

    cols = [
        'Instituicao_Lotacao', 'Camas_Totais_Media',
        'Taxa_Ocupacao_Media_Pct', 'Capacidade_Real',
    ]
    final = merged[cols].copy()
    final.columns = [
        'ULS', 'Camas Totais (média 2025)',
        'Taxa Ocupação (%)', 'Expectativa de capacidade utilizada',
    ]
    final = final.sort_values('ULS').reset_index(drop=True)

    return final, sem_ocup


# ══════════════════════════════════════════════════════════════════════════
# 4. EXCEL WRITER
# ══════════════════════════════════════════════════════════════════════════

def fill(hex_color): return PatternFill('solid', fgColor=hex_color)

def thin_border():
    s = Side(style='thin', color='AAAAAA')
    return Border(left=s, right=s, top=s, bottom=s)

def center(): return Alignment(horizontal='center', vertical='center', wrap_text=True)
def left():   return Alignment(horizontal='left', vertical='center', wrap_text=True)


def write_df_to_ws(ws, df, start_row=1, start_col=1, header_bg=C_HEADER):
    cols = list(df.columns)
    for ci, col in enumerate(cols, start_col):
        cell = ws.cell(row=start_row, column=ci, value=col)
        cell.font = Font(bold=True, color='FFFFFF', name='Calibri', size=11)
        cell.fill = fill(header_bg)
        cell.alignment = center()
        cell.border = thin_border()
    for ri, (_, row_data) in enumerate(df.iterrows(), start_row + 1):
        alt = (ri - start_row) % 2 == 0
        for ci, val in enumerate(row_data, start_col):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.fill = fill(C_ALT if alt else C_WHITE)
            cell.alignment = left() if isinstance(val, str) else center()
            cell.border = thin_border()
    for ci, col in enumerate(cols, start_col):
        max_len = max(len(str(col)), df[col].astype(str).str.len().max() if len(df) > 0 else 10)
        ws.column_dimensions[get_column_letter(ci)].width = min(max_len + 3, 50)
    return start_row + len(df) + 2


def add_note(ws, row, col, text, n_cols=1):
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = Font(italic=True, size=9, color='444444')
    cell.alignment = left()
    if n_cols > 1:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + n_cols - 1)
    ws.row_dimensions[row].height = 30
    return row + 1


def build_excel(final, sem_ocup):
    wb = Workbook()
    wb.remove(wb.active)

    ws1 = wb.create_sheet('1. Oferta por ULS')
    ws1.sheet_view.showGridLines = False
    r = add_note(ws1, 1, 1,
        "Capacidade Real = Camas Totais × Taxa de Ocupação. "
        "Representa o número médio de camas efetivamente ocupadas por dia, em 2025.", 4)
    r = add_note(ws1, r, 1,
        "Camas Totais = soma de Camas Cirúrgicas + Médicas + Neutras + Outras, média mensal de 2025. "
        "Taxa de Ocupação = média mensal da taxa de ocupação reportada para a ULS em 2025.", 4)
    write_df_to_ws(ws1, final, start_row=r, header_bg=C_SUBHDR)

    wb.save(OUTPUT_FILE)
    print(f'✅ Excel da Oferta guardado em: {OUTPUT_FILE}')


# ══════════════════════════════════════════════════════════════════════════
# MAIN
# ══════════════════════════════════════════════════════════════════════════

def main():
    print('🔄 A calcular oferta (capacidade real) por ULS...')
    final, sem_ocup = build_oferta()

    print(f'🔄 A construir Excel ({len(final)} ULS)...')
    build_excel(final, sem_ocup)

    if sem_ocup:
        print(f'⚠️  {len(sem_ocup)} ULS sem Taxa de Ocupação:')
        for u in sem_ocup:
            print(f'   • {u}')
    else:
        print('✅ Todas as ULS têm Taxa de Ocupação correspondente.')


if __name__ == '__main__':
    main()