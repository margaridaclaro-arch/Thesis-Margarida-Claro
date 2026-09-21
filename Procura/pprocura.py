"""
Modelo de Procura Hospitalar por Freguesia × Faixa Etária
=====================================================================
Fórmulas:
  Probabilidade nacional     = Internamentos Nacionais (faixa X, doença Y) / Pop. Nacional faixa X
  Internamentos Esperados    = Probabilidade × Pop. Freguesia na Faixa
  Dias Necessários           = Internamentos Esperados × Demora Média Nacional
  Camas Necessárias          = Dias Necessários / 365

Ficheiros necessários (mesma pasta que o script):
  - freguesias_limpo.xlsx
  - morbilidade_mortalidade_hospit.xlsx
"""

import pandas as pd
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR      = os.path.dirname(os.path.abspath(__file__))
PARISHES_FILE = os.path.join(BASE_DIR, 'freguesias_limpo.xlsx')
MORB_FILE     = os.path.join(BASE_DIR, 'morbilidade_mortalidade_hospit.xlsx')
OUTPUT_FILE   = os.path.join(BASE_DIR, 'Excel_Procura_Portugal.xlsx')

FAIXAS_ORDEM = ['[0-1[', '[1-5[', '[5-15[', '[15-25[', '[25-45[', '[45-65[', '[65-120[']
FAIXA_TO_COL = {
    '[0-1[':    'Pop_0_1',
    '[1-5[':    'Pop_1_5',
    '[5-15[':   'Pop_5_15',
    '[15-25[':  'Pop_15_25',
    '[25-45[':  'Pop_25_45',
    '[45-65[':  'Pop_45_65',
    '[65-120[': 'Pop_65plus',
}
col_to_faixa = {v: k for k, v in FAIXA_TO_COL.items()}


# ══════════════════════════════════════════════════════════════════════════
# 1. POPULAÇÃO POR FREGUESIA
# ══════════════════════════════════════════════════════════════════════════

def load_populacao():
    df = pd.read_excel(PARISHES_FILE, sheet_name='Sheet1', header=0)
    df.columns = [
        'Distrito', 'Concelho', 'Freguesia', 'Freguesia_Concelho', 'Pop_Total',
        'Pop_0_1', 'Pop_1_5', 'Pop_5_15', 'Pop_15_25', 'Pop_25_45', 'Pop_45_65', 'Pop_65plus',
        'Lat_Freguesia', 'Long_Freguesia', 'ULS', 'Hospital', 'Lat_Hospital', 'Long_Hospital'
    ]
    df = df.dropna(subset=['ULS'])
    for c in list(FAIXA_TO_COL.values()) + ['Pop_Total']:
        df[c] = pd.to_numeric(df[c], errors='coerce').fillna(0)
    return df


# ══════════════════════════════════════════════════════════════════════════
# 2. PROBABILIDADES NACIONAIS POR FAIXA × DOENÇA
# ══════════════════════════════════════════════════════════════════════════

def load_probabilidades_nacionais(pop_df):
    df = pd.read_excel(MORB_FILE, sheet_name='Feuil1', header=0)
    df.columns = [
        'Periodo', 'Regiao', 'Instituicao', 'Cod_Cap', 'Descricao',
        'Faixa_Etaria', 'Sexo', 'Internamentos', 'Dias_Internamento', 'Ambulatorio', 'Obitos'
    ]
    df['Internamentos']     = pd.to_numeric(df['Internamentos'],     errors='coerce').fillna(0)
    df['Dias_Internamento'] = pd.to_numeric(df['Dias_Internamento'], errors='coerce').fillna(0)
    df = df[df['Instituicao'].str.contains('Unidade Local de Saúde', na=False)]

    # CORRECAO: o ficheiro de origem contem linhas repetidas para a mesma
    # combinacao (Periodo, Regiao, Instituicao, Cod_Cap, Faixa_Etaria, Sexo)
    # com valores ligeiramente diferentes (varias extracoes/revisoes do
    # mesmo relatorio). Sem esta deduplicacao, o groupby().sum() seguinte
    # conta a mesma admissao varias vezes, inflacionando os internamentos
    # nacionais em ~33%.
    # Usa-se o MAXIMO (nao a media): a documentacao oficial do dataset (SNS
    # Transparencia) explica que a codificacao clinica e feita apos a alta
    # e pode atrasar-se, causando SUBCONTAGEM temporaria nalgumas extracoes
    # - nao ha indicacao de correcoes para baixo. Logo, entre varias
    # extracoes do mesmo periodo, a de maior valor e a mais completa/fiavel.
    chave_unica = ['Periodo', 'Regiao', 'Instituicao', 'Cod_Cap', 'Faixa_Etaria', 'Sexo']
    df = df.groupby(chave_unica + ['Descricao'], as_index=False).agg(
        Internamentos=('Internamentos', 'max'),
        Dias_Internamento=('Dias_Internamento', 'max'),
    )

    nacional = df.groupby(['Faixa_Etaria', 'Cod_Cap', 'Descricao']).agg(
        Internamentos_Nac=('Internamentos',     'sum'),
        Dias_Nac         =('Dias_Internamento', 'sum'),
    ).reset_index()
    nacional = nacional[nacional['Internamentos_Nac'] > 0].copy()

    pop_nac = pop_df[list(FAIXA_TO_COL.values())].sum()
    faixa_to_popnac = {faixa: pop_nac[col] for col, faixa in col_to_faixa.items()}
    nacional['Pop_Nacional_Faixa'] = nacional['Faixa_Etaria'].map(faixa_to_popnac)
    nacional['Probabilidade'] = (nacional['Internamentos_Nac'] / nacional['Pop_Nacional_Faixa']).round(8)
    nacional['Demora_Media']  = (nacional['Dias_Nac'] / nacional['Internamentos_Nac']).round(2)

    return nacional


# ══════════════════════════════════════════════════════════════════════════
# 3. CÁLCULO DA PROCURA
# ══════════════════════════════════════════════════════════════════════════

def build_procura(pop_df, nacional):
    pop_long = pop_df[
        ['Distrito', 'Concelho', 'Freguesia', 'Pop_Total'] + list(FAIXA_TO_COL.values())
    ].melt(
        id_vars=['Distrito', 'Concelho', 'Freguesia', 'Pop_Total'],
        value_vars=list(FAIXA_TO_COL.values()),
        var_name='Pop_Col', value_name='Pop_Faixa'
    )
    pop_long['Faixa_Etaria'] = pop_long['Pop_Col'].map(col_to_faixa)
    pop_long = pop_long.drop(columns='Pop_Col')

    merged = pop_long.merge(nacional, on='Faixa_Etaria', how='left')
    merged['Internamentos_Esperados'] = (merged['Probabilidade'] * merged['Pop_Faixa']).round(2)
    merged['Dias_Necessarios']        = (merged['Internamentos_Esperados'] * merged['Demora_Media']).round(2)
    merged['Camas_Necessarias']       = (merged['Dias_Necessarios'] / 365).round(4)

    merged['_ordem'] = merged['Faixa_Etaria'].apply(
        lambda f: FAIXAS_ORDEM.index(f) if f in FAIXAS_ORDEM else 99
    )
    merged = merged.sort_values(
        ['Distrito', 'Concelho', 'Freguesia', '_ordem', 'Descricao']
    ).drop(columns='_ordem').reset_index(drop=True)

    return merged


def build_agregado_faixa(merged):
    agg = merged.groupby(['Distrito', 'Concelho', 'Freguesia', 'Faixa_Etaria']).agg(
        Pop_Faixa              =('Pop_Faixa',               'first'),
        Internamentos_Esperados=('Internamentos_Esperados',  'sum'),
        Dias_Necessarios       =('Dias_Necessarios',         'sum'),
        Camas_Necessarias      =('Camas_Necessarias',        'sum'),
    ).reset_index()
    agg['Demora_Media_Pond'] = (agg['Dias_Necessarios'] / agg['Internamentos_Esperados']).round(2)
    agg['_ordem'] = agg['Faixa_Etaria'].apply(
        lambda f: FAIXAS_ORDEM.index(f) if f in FAIXAS_ORDEM else 99
    )
    return agg.sort_values(
        ['Distrito', 'Concelho', 'Freguesia', '_ordem']
    ).drop(columns='_ordem').reset_index(drop=True)


# ══════════════════════════════════════════════════════════════════════════
# 4. EXCEL
# ══════════════════════════════════════════════════════════════════════════

def fill(h):  return PatternFill('solid', fgColor=h)
def tb():
    s = Side(style='thin', color='AAAAAA')
    return Border(left=s, right=s, top=s, bottom=s)
def cen(): return Alignment(horizontal='center', vertical='center', wrap_text=True)
def lft(): return Alignment(horizontal='left',   vertical='center', wrap_text=True)


def write_df_formatted(ws, df, sr=1, sc=1, hbg='1F3864', formula_cols=None):
    """Escrita com formatação completa — para folhas pequenas."""
    formula_cols = formula_cols or {}
    for ci, col in enumerate(df.columns, sc):
        c = ws.cell(row=sr, column=ci, value=col)
        c.font      = Font(bold=True, color='FFFFFF', name='Calibri', size=11)
        c.fill      = fill(formula_cols.get(col, hbg))
        c.alignment = cen()
        c.border    = tb()
    for ri, (_, row) in enumerate(df.iterrows(), sr + 1):
        alt = (ri - sr) % 2 == 0
        for ci, val in enumerate(row, sc):
            c           = ws.cell(row=ri, column=ci, value=val)
            c.fill      = fill('D9E1F2' if alt else 'FFFFFF')
            c.alignment = lft() if isinstance(val, str) else cen()
            c.border    = tb()
    for ci, col in enumerate(df.columns, sc):
        ml = max(len(str(col)), df[col].astype(str).str.len().max() if len(df) > 0 else 10)
        ws.column_dimensions[get_column_letter(ci)].width = min(ml + 3, 55)
    return sr + len(df) + 2


def write_df_fast(ws, df, sr=1, sc=1, hbg='1F3864', formula_cols=None):
    """Escrita rápida: só cabeçalho formatado, dados sem formatação — para folhas grandes."""
    formula_cols = formula_cols or {}
    # Cabeçalho formatado
    for ci, col in enumerate(df.columns, sc):
        c = ws.cell(row=sr, column=ci, value=col)
        c.font      = Font(bold=True, color='FFFFFF', name='Calibri', size=11)
        c.fill      = fill(formula_cols.get(col, hbg))
        c.alignment = cen()
        c.border    = tb()
    # Dados sem formatação (append_rows é o método mais rápido disponível)
    for row in df.itertuples(index=False):
        ws.append(list(row))
    # Largura das colunas
    for ci, col in enumerate(df.columns, sc):
        ml = max(len(str(col)), df[col].astype(str).str.len().max() if len(df) > 0 else 10)
        ws.column_dimensions[get_column_letter(ci)].width = min(ml + 3, 55)
    return sr + len(df) + 2


def note(ws, r, c, t, n=1, bold=False):
    cell           = ws.cell(row=r, column=c, value=t)
    cell.font      = Font(italic=not bold, bold=bold, size=9, color='444444')
    cell.alignment = lft()
    if n > 1:
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + n - 1)
    ws.row_dimensions[r].height = 28
    return r + 1


def build_excel(nacional, merged, agregado_faixa):
    wb = Workbook()
    wb.remove(wb.active)

    # ── Folha 0: Probabilidades Nacionais (formatada) ──────────────────────
    ws0 = wb.create_sheet('0. Probabilidades Nacionais')
    ws0.sheet_view.showGridLines = False
    r = note(ws0, 1, 1, 'Fórmulas:', 8, bold=True)
    r = note(ws0, r, 1, '  Probabilidade  =  Internamentos Nacionais / Pop. Nacional da Faixa', 8)
    r = note(ws0, r, 1, '  Demora Média Nacional  =  Dias Nacionais de Internamento / Internamentos Nacionais'
             '  →  usada nas folhas seguintes para calcular Dias Necessários por freguesia.', 8)

    df0 = nacional[[
        'Faixa_Etaria', 'Cod_Cap', 'Descricao',
        'Pop_Nacional_Faixa', 'Internamentos_Nac', 'Dias_Nac',
        'Demora_Media', 'Probabilidade'
    ]].copy()
    df0['_o'] = df0['Faixa_Etaria'].apply(lambda f: FAIXAS_ORDEM.index(f) if f in FAIXAS_ORDEM else 99)
    df0 = df0.sort_values(['_o', 'Descricao']).drop(columns='_o')
    df0.columns = [
        'Faixa Etária', 'Cód.', 'Diagnóstico',
        'Pop. Nacional da Faixa', 'Internamentos Nacionais', 'Dias Nacionais de Internamento',
        'Demora Média Nacional (dias)  =  Dias Nac. / Internamentos Nac.',
        'Probabilidade  =  Internamentos Nac. / Pop. Nacional da Faixa',
    ]
    write_df_formatted(ws0, df0, sr=r, hbg='006400', formula_cols={
        'Demora Média Nacional (dias)  =  Dias Nac. / Internamentos Nac.': '1A5276',
        'Probabilidade  =  Internamentos Nac. / Pop. Nacional da Faixa':   '1A5276',
    })
    print('   Folha 0 OK')

    # ── Folha 1: Detalhe por Doença (rápida, sem formatação nos dados) ─────
    ws1 = wb.create_sheet('1. Detalhe por Doenca')
    ws1.sheet_view.showGridLines = False
    r1 = note(ws1, 1, 1, 'Detalhe por freguesia × faixa × doença. Cabeçalho formatado; dados sem formatação para rapidez de escrita.', 10, bold=True)
    r1 = note(ws1, r1, 1, '  Internamentos Esperados  =  Probabilidade Nacional (faixa × doença)  ×  Pop. da Faixa na Freguesia', 10)
    r1 = note(ws1, r1, 1, '  Dias Necessários          =  Internamentos Esperados  ×  Demora Média Nacional', 10)
    r1 = note(ws1, r1, 1, '  Camas Necessárias         =  Dias Necessários  /  365   (aditivas: Σ camas por doença = camas totais)', 10)

    df1 = merged[[
        'Distrito', 'Concelho', 'Freguesia', 'Faixa_Etaria',
        'Cod_Cap', 'Descricao', 'Pop_Faixa',
        'Probabilidade', 'Internamentos_Esperados',
        'Demora_Media', 'Dias_Necessarios', 'Camas_Necessarias'
    ]].copy()
    df1.columns = [
        'Distrito', 'Concelho', 'Freguesia', 'Faixa Etária',
        'Cód.', 'Diagnóstico', 'Pop. da Faixa',
        'Probabilidade',
        'Internamentos Esperados  =  Prob × Pop. Faixa',
        'Demora Média Nacional (dias)',
        'Dias Necessários  =  Internamentos × Demora Média',
        'Camas Necessárias  =  Dias / 365',
    ]
    write_df_fast(ws1, df1, sr=r1, hbg='1F3864', formula_cols={
        'Internamentos Esperados  =  Prob × Pop. Faixa':     '1A5276',
        'Dias Necessários  =  Internamentos × Demora Média': '1A5276',
        'Camas Necessárias  =  Dias / 365':                  '1A5276',
    })
    print('   Folha 1 OK')

    # ── Folha 2: Agregado por Freguesia × Faixa (formatada) ───────────────
    ws2 = wb.create_sheet('2. Agregado por Faixa')
    ws2.sheet_view.showGridLines = False
    r2 = note(ws2, 1, 1, 'Soma de todas as doenças por freguesia × faixa etária. Totais iguais aos da folha 1 agregados.', 9, bold=True)
    r2 = note(ws2, r2, 1, '  Internamentos Esperados  =  Σ doenças (Prob × Pop. Faixa)', 9)
    r2 = note(ws2, r2, 1, '  Dias Necessários          =  Internamentos Esperados  ×  Demora Média Ponderada', 9)
    r2 = note(ws2, r2, 1, '  Camas Necessárias         =  Dias Necessários  /  365', 9)

    df2 = agregado_faixa[[
        'Distrito', 'Concelho', 'Freguesia', 'Faixa_Etaria', 'Pop_Faixa',
        'Internamentos_Esperados', 'Demora_Media_Pond', 'Dias_Necessarios', 'Camas_Necessarias'
    ]].copy()
    df2.columns = [
        'Distrito', 'Concelho', 'Freguesia', 'Faixa Etária', 'Pop. da Faixa',
        'Internamentos Esperados  =  Σ(Prob × Pop. Faixa)',
        'Demora Média Pond. (dias)',
        'Dias Necessários  =  Internamentos × Demora Média',
        'Camas Necessárias  =  Dias / 365',
    ]
    write_df_formatted(ws2, df2, sr=r2, hbg='4A148C', formula_cols={
        'Internamentos Esperados  =  Σ(Prob × Pop. Faixa)':       '1A5276',
        'Dias Necessários  =  Internamentos × Demora Média':      '1A5276',
        'Camas Necessárias  =  Dias / 365':                       '1A5276',
    })
    print('   Folha 2 OK')

    wb.save(OUTPUT_FILE)
    print(f'✅ Excel guardado em: {OUTPUT_FILE}')


# ══════════════════════════════════════════════════════════════════════════
# MAIN
# ══════════════════════════════════════════════════════════════════════════

def main():
    if not os.path.exists(PARISHES_FILE) or not os.path.exists(MORB_FILE):
        print(f'\n❌ ERRO: coloque os ficheiros Excel na mesma pasta que este script:\n   {BASE_DIR}')
        return

    print('🔄 A carregar população por freguesia...')
    pop_df = load_populacao()
    print(f'   {len(pop_df)} freguesias | {pop_df["ULS"].nunique()} ULS')

    print('🔄 A calcular probabilidades nacionais por faixa × doença...')
    nacional = load_probabilidades_nacionais(pop_df)
    print(f'   {len(nacional)} combinações faixa × doença com internamentos > 0')

    print('🔄 A calcular internamentos esperados por freguesia × faixa × doença...')
    merged = build_procura(pop_df, nacional)
    print(f'   {len(merged):,} linhas de detalhe')

    print('🔄 A agregar por freguesia × faixa...')
    agregado_faixa = build_agregado_faixa(merged)
    print(f'   {len(agregado_faixa):,} linhas agregadas')

    print('🔄 A construir Excel...')
    build_excel(nacional, merged, agregado_faixa)


if __name__ == '__main__':
    main()