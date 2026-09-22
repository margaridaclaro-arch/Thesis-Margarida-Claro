#!/usr/bin/env python3
"""
GERAR DASHBOARD INTERATIVO — MODELO 2 (Localização Ótima de Novo Hospital no Algarve)
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
import os

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
INPUT_FILE = os.path.join(SCRIPT_DIR, "Modelo2_Localizacao_Algarve.xlsx")
OUTPUT_FILE = os.path.join(SCRIPT_DIR, "Modelo2_Localizacao_Algarve_Dashboard.xlsx")

# ── Estilo (verde — cor de destaque do Modelo 2) ────────────────────────────
VERDE       = "2E9E6B"
VERDE_ESC   = "1F6E4B"
VERDE_CLARO = "DCEFE6"
LARANJA     = "E07A2C"
AZUL        = "2E75B6"
VERMELHO    = "C0392B"
CINZA       = "F2F2F2"
FONT_NAME   = "Arial"

HEADER_FILL = PatternFill(start_color=VERDE, end_color=VERDE, fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
KPI_FILL    = PatternFill(start_color=VERDE_CLARO, end_color=VERDE_CLARO, fill_type="solid")
TITLE_FONT  = Font(name=FONT_NAME, bold=True, size=18, color=VERDE_ESC)
SUB_FONT    = Font(name=FONT_NAME, italic=True, size=10, color="666666")
LABEL_FONT  = Font(name=FONT_NAME, bold=True, size=9, color="666666")
KPI_FONT    = Font(name=FONT_NAME, bold=True, size=18, color=VERDE_ESC)
SECTION_FONT= Font(name=FONT_NAME, bold=True, size=12, color="FFFFFF")
SECTION_FILL= PatternFill(start_color=VERDE, end_color=VERDE, fill_type="solid")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CONCELHOS = ['Albufeira', 'Alcoutim', 'Aljezur', 'Castro Marim', 'Faro', 'Lagoa', 'Lagos',
             'Loulé', 'Monchique', 'Olhão', 'Portimão', 'Silves', 'São Brás de Alportel',
             'Tavira', 'Vila Real de Santo António', 'Vila do Bispo']

N_FREG_LAST_ROW = 69   # folhas 2 e 3: dados em linhas 3..69 (67 freguesias)

S_IMPACTO = "'3. Impacto_Melhor_Local'"   # A=Concelho B=Freguesia C=Populacao D=Camas_Necess.
                                           # E=tempo_atual F=tempo_novo G=poupanca H=fracao I=usa_novo
S_RANK    = "'2. Ranking_Candidatos'"     # A=Ranking B=Concelho C=Freguesia D=Populacao
                                           # E=Lat F=Long G=tempo_medio_se_aqui H=defice I=custo


def style_section_title(ws, cell_range, text):
    ws.merge_cells(cell_range)
    top_left = cell_range.split(":")[0]
    ws[top_left] = text
    ws[top_left].font = SECTION_FONT
    ws[top_left].fill = SECTION_FILL
    ws[top_left].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    row = int(''.join(filter(str.isdigit, top_left)))
    ws.row_dimensions[row].height = 22


def kpi_card(ws, cell_range, label, formula, font=KPI_FONT, num_fmt=None):
    ws.merge_cells(cell_range)
    top_left = cell_range.split(":")[0]
    col = ''.join(filter(str.isalpha, top_left))
    row = int(''.join(filter(str.isdigit, top_left)))
    lbl_row = row - 1
    ws.cell(row=lbl_row, column=openpyxl.utils.column_index_from_string(col),
            value=label).font = LABEL_FONT
    c = ws[top_left]
    c.value = formula
    c.font = font
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.fill = KPI_FILL
    if num_fmt:
        c.number_format = num_fmt
    for r in ws[cell_range]:
        for cell in r:
            cell.fill = KPI_FILL
            cell.border = BORDER


def main():
    print("A carregar workbook do modelo...")
    wb = openpyxl.load_workbook(INPUT_FILE)

    if "0. Dashboard" in wb.sheetnames:
        del wb["0. Dashboard"]
    ws = wb.create_sheet("0. Dashboard", 0)
    ws.sheet_view.showGridLines = False

    widths = {"A": 3, "B": 19, "C": 19, "D": 19, "E": 19, "F": 19, "G": 19,
              "H": 19, "I": 19, "J": 19, "K": 4}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # ── Título ──────────────────────────────────────────────────────────
    ws.merge_cells("B2:J2")
    ws["B2"] = "DASHBOARD INTERATIVO — NOVO HOSPITAL NO ALGARVE (MODELO 2)"
    ws["B2"].font = TITLE_FONT
    ws.merge_cells("B3:J3")
    ws["B3"] = ("Em que freguesia deveria ficar um novo hospital de 742 camas para minimizar o tempo de "
                "trajeto ponderado por população? — 67 freguesias candidatas do Algarve, capacidade real")
    ws["B3"].font = SUB_FONT
    ws.row_dimensions[2].height = 26

    # ── Banner: melhor localização ─────────────────────────────────────
    ws.merge_cells("B5:J6")
    ws["B5"] = '🏆  MELHOR LOCALIZAÇÃO: PORTIMÃO (CONCELHO DE PORTIMÃO)  —  ="lat/long: "&\'1. Resumo\'!C11'
    # a linha acima é substituída de seguida por uma fórmula real (concatenação)
    ws["B5"] = '="🏆  MELHOR LOCALIZAÇÃO: "&UPPER(\'1. Resumo\'!C9)&"  (concelho de "&\'1. Resumo\'!C10&")   —   coordenadas: "&\'1. Resumo\'!C11'
    ws["B5"].font = Font(name=FONT_NAME, bold=True, size=14, color="FFFFFF")
    ws["B5"].fill = PatternFill(start_color=VERDE_ESC, end_color=VERDE_ESC, fill_type="solid")
    ws["B5"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[5].height = 22
    ws.row_dimensions[6].height = 22

    # ── KPIs nacionais (ligados à folha 1. Resumo) ─────────────────────
    row_lbl, row_val = 8, 9
    kpi_card(ws, f"B{row_val}:B{row_val}", "Nº FREGUESIAS", "='1. Resumo'!C5", num_fmt="#,##0")
    kpi_card(ws, f"C{row_val}:C{row_val}", "POPULAÇÃO DO ALGARVE", "='1. Resumo'!C6", num_fmt="#,##0")
    kpi_card(ws, f"D{row_val}:D{row_val}", "TEMPO MÉDIO ATUAL", "='1. Resumo'!C13", num_fmt='0.0" min"')
    kpi_card(ws, f"E{row_val}:E{row_val}", "TEMPO MÉDIO C/ NOVO HOSP.", "='1. Resumo'!C14", num_fmt='0.0" min"')
    kpi_card(ws, f"F{row_val}:F{row_val}", "REDUÇÃO DE TEMPO", "='1. Resumo'!C15", num_fmt='0.0" min"')
    kpi_card(ws, f"G{row_val}:G{row_val}", "% CRÍTICAS (>60min) ATUAL", "='1. Resumo'!C17/100", num_fmt="0.0%")
    kpi_card(ws, f"H{row_val}:H{row_val}", "% CRÍTICAS (>60min) ÓTIMO", "='1. Resumo'!C18/100", num_fmt="0.0%")
    kpi_card(ws, f"I{row_val}:I{row_val}", "FREG. QUE USAM NOVO HOSP.",
             "='1. Resumo'!C20&\" / \"&'1. Resumo'!C5", font=Font(name=FONT_NAME, bold=True, size=15, color=VERDE_ESC))
    kpi_card(ws, f"J{row_val}:J{row_val}", "DÉFICE DE CAMAS (ÓTIMO)", "='1. Resumo'!C29", num_fmt='0.0" camas"')
    for c in "BCDEFGHIJ":
        ws.row_dimensions[row_val].height = 30

    # ── Filtro interativo por Concelho ─────────────────────────────────
    style_section_title(ws, "B12:J12", "🔎  FILTRO INTERATIVO — RESUMO POR CONCELHO SELECIONADO")

    ws["B14"] = "Escolher Concelho:"
    ws["B14"].font = Font(name=FONT_NAME, bold=True, size=11)
    ws["C14"] = "Portimão"
    ws["C14"].font = Font(name=FONT_NAME, bold=True, size=11, color=VERMELHO)
    ws["C14"].fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    ws["C14"].alignment = Alignment(horizontal="center")
    ws["C14"].border = BORDER

    ws["N1"] = "lista_concelhos"
    for i, cne in enumerate(CONCELHOS, start=2):
        ws.cell(row=i, column=14, value=cne)
    ws.column_dimensions["N"].width = 2

    dv = DataValidation(type="list", formula1="=$N$2:$N$17", allow_blank=False, showDropDown=False)
    dv.error = "Escolha um concelho da lista."
    dv.errorTitle = "Concelho inválido"
    ws.add_data_validation(dv)
    dv.add(ws["C14"])

    filtro_labels = ["Nº Freguesias", "População total", "Tempo médio ATUAL (pond.)",
                      "Tempo médio C/ NOVO HOSP. (pond.)", "Poupança média", "% que usa o novo hospital"]
    filtro_formulas = [
        f'=COUNTIF({S_IMPACTO}!A3:A{N_FREG_LAST_ROW},$C$14)',
        f'=SUMIF({S_IMPACTO}!A3:A{N_FREG_LAST_ROW},$C$14,{S_IMPACTO}!C3:C{N_FREG_LAST_ROW})',
        f'=SUMPRODUCT(({S_IMPACTO}!$A3:A{N_FREG_LAST_ROW}=$C$14)*{S_IMPACTO}!C3:C{N_FREG_LAST_ROW}*{S_IMPACTO}!E3:E{N_FREG_LAST_ROW})'
        f'/SUMPRODUCT(({S_IMPACTO}!$A3:A{N_FREG_LAST_ROW}=$C$14)*{S_IMPACTO}!C3:C{N_FREG_LAST_ROW})',
        f'=SUMPRODUCT(({S_IMPACTO}!$A3:A{N_FREG_LAST_ROW}=$C$14)*{S_IMPACTO}!C3:C{N_FREG_LAST_ROW}*{S_IMPACTO}!F3:F{N_FREG_LAST_ROW})'
        f'/SUMPRODUCT(({S_IMPACTO}!$A3:A{N_FREG_LAST_ROW}=$C$14)*{S_IMPACTO}!C3:C{N_FREG_LAST_ROW})',
        f'=SUMIF({S_IMPACTO}!A3:A{N_FREG_LAST_ROW},$C$14,{S_IMPACTO}!G3:G{N_FREG_LAST_ROW})/COUNTIF({S_IMPACTO}!A3:A{N_FREG_LAST_ROW},$C$14)',
        f'=COUNTIFS({S_IMPACTO}!A3:A{N_FREG_LAST_ROW},$C$14,{S_IMPACTO}!I3:I{N_FREG_LAST_ROW},"Sim")'
        f'/COUNTIF({S_IMPACTO}!A3:A{N_FREG_LAST_ROW},$C$14)',
    ]
    fmts = ["#,##0", "#,##0", '0.0" min"', '0.0" min"', '0.0" min"', "0.0%"]
    col0 = 4
    for i, (lbl, formula, fmt) in enumerate(zip(filtro_labels, filtro_formulas, fmts)):
        col = col0 + i
        ws.cell(row=13, column=col, value=lbl).font = LABEL_FONT
        cell = ws.cell(row=14, column=col, value=formula)
        cell.font = Font(name=FONT_NAME, bold=True, size=12, color=VERDE_ESC)
        cell.number_format = fmt
        cell.fill = KPI_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER
    ws.row_dimensions[14].height = 24

    # ── Tabela: Análise por Concelho (fórmulas) ─────────────────────────
    style_section_title(ws, "B17:F17", "📊  ANÁLISE POR CONCELHO")
    hdr_row = 18
    headers = ["Concelho", "Nº Freguesias", "Tempo Médio ATUAL (min)",
               "Tempo Médio C/ NOVO HOSP. (min)", "% Usa Novo Hospital"]
    for j, h in enumerate(headers, 2):
        c = ws.cell(row=hdr_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[hdr_row].height = 30

    conc_start_row = hdr_row + 1
    for i, cne in enumerate(CONCELHOS):
        r = conc_start_row + i
        ws.cell(row=r, column=2, value=cne).border = BORDER
        ws.cell(row=r, column=3,
                value=f'=COUNTIF({S_IMPACTO}!$A$3:$A${N_FREG_LAST_ROW},$B{r})').border = BORDER
        ws.cell(row=r, column=4,
                value=(f'=IFERROR(ROUND(SUMPRODUCT(({S_IMPACTO}!$A$3:$A${N_FREG_LAST_ROW}=$B{r})*{S_IMPACTO}!$C$3:$C${N_FREG_LAST_ROW}*{S_IMPACTO}!$E$3:$E${N_FREG_LAST_ROW})'
                       f'/SUMPRODUCT(({S_IMPACTO}!$A$3:$A${N_FREG_LAST_ROW}=$B{r})*{S_IMPACTO}!$C$3:$C${N_FREG_LAST_ROW}),1),0)')
                ).border = BORDER
        ws.cell(row=r, column=5,
                value=(f'=IFERROR(ROUND(SUMPRODUCT(({S_IMPACTO}!$A$3:$A${N_FREG_LAST_ROW}=$B{r})*{S_IMPACTO}!$C$3:$C${N_FREG_LAST_ROW}*{S_IMPACTO}!$F$3:$F${N_FREG_LAST_ROW})'
                       f'/SUMPRODUCT(({S_IMPACTO}!$A$3:$A${N_FREG_LAST_ROW}=$B{r})*{S_IMPACTO}!$C$3:$C${N_FREG_LAST_ROW}),1),0)')
                ).border = BORDER
        ws.cell(row=r, column=6,
                value=(f'=COUNTIFS({S_IMPACTO}!$A$3:$A${N_FREG_LAST_ROW},$B{r},{S_IMPACTO}!$I$3:$I${N_FREG_LAST_ROW},"Sim")'
                       f'/COUNTIF({S_IMPACTO}!$A$3:$A${N_FREG_LAST_ROW},$B{r})')
                ).number_format = "0.0%"
        ws.cell(row=r, column=6).border = BORDER
        for col in (2, 3, 4, 5, 6):
            ws.cell(row=r, column=col).alignment = Alignment(horizontal="center")
            if (i % 2) == 1:
                ws.cell(row=r, column=col).fill = PatternFill(start_color=CINZA, end_color=CINZA, fill_type="solid")
    conc_end_row = conc_start_row + len(CONCELHOS) - 1

    # ── Tabela: Top 10 candidatos a melhor localização ──────────────────
    style_section_title(ws, "H17:J17", "🏆  TOP 10 CANDIDATOS (menor custo = melhor)")
    ws.cell(row=hdr_row, column=8, value="Freguesia (Concelho)").font = HEADER_FONT
    ws.cell(row=hdr_row, column=8).fill = HEADER_FILL
    ws.cell(row=hdr_row, column=8).border = BORDER
    ws.cell(row=hdr_row, column=9, value="Tempo Médio (min)").font = HEADER_FONT
    ws.cell(row=hdr_row, column=9).fill = HEADER_FILL
    ws.cell(row=hdr_row, column=9).alignment = Alignment(horizontal="center", wrap_text=True)
    ws.cell(row=hdr_row, column=9).border = BORDER
    ws.merge_cells(start_row=hdr_row, start_column=9, end_row=hdr_row, end_column=10)

    top10_start = hdr_row + 1
    for i in range(10):
        src_row = 3 + i   # sheet2 já vem ordenado por custo_ponderado ascendente (melhor primeiro)
        r = top10_start + i
        ws.cell(row=r, column=8,
                value=f'={S_RANK}!C{src_row}&" ("&{S_RANK}!B{src_row}&")"').border = BORDER
        val_cell = ws.cell(row=r, column=9, value=f'={S_RANK}!G{src_row}')
        ws.merge_cells(start_row=r, start_column=9, end_row=r, end_column=10)
        val_cell.number_format = "0.00"
        val_cell.alignment = Alignment(horizontal="center")
        val_cell.border = BORDER
        ws.cell(row=r, column=8).alignment = Alignment(horizontal="left")
        if i % 2 == 1:
            for col in (8, 9, 10):
                ws.cell(row=r, column=col).fill = PatternFill(start_color=CINZA, end_color=CINZA, fill_type="solid")
    top10_end = top10_start + 9

    # ── Tabela: Top 15 freguesias com maior poupança de tempo ──────────
    top15_title_row = conc_end_row + 3
    style_section_title(ws, f"B{top15_title_row}:F{top15_title_row}",
                         "⏱️  TOP 15 FREGUESIAS COM MAIOR POUPANÇA DE TEMPO (folha 3, já ordenada)")
    top15_hdr = top15_title_row + 1
    headers15 = ["Freguesia", "Concelho", "Tempo Atual (min)", "Tempo C/ Novo Hosp. (min)", "Poupança (min)"]
    for j, h in enumerate(headers15, 2):
        c = ws.cell(row=top15_hdr, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[top15_hdr].height = 28
    top15_start = top15_hdr + 1
    for i in range(15):
        src_row = 3 + i  # folha 3 já vem ordenada por poupanca_min desc
        r = top15_start + i
        ws.cell(row=r, column=2, value=f"={S_IMPACTO}!B{src_row}").border = BORDER
        ws.cell(row=r, column=3, value=f"={S_IMPACTO}!A{src_row}").border = BORDER
        ws.cell(row=r, column=4, value=f"={S_IMPACTO}!E{src_row}").border = BORDER
        ws.cell(row=r, column=4).number_format = "0.0"
        ws.cell(row=r, column=5, value=f"={S_IMPACTO}!F{src_row}").border = BORDER
        ws.cell(row=r, column=5).number_format = "0.0"
        ws.cell(row=r, column=6, value=f"={S_IMPACTO}!G{src_row}").border = BORDER
        ws.cell(row=r, column=6).number_format = "0.0"
        for col in (2, 3, 4, 5, 6):
            ws.cell(row=r, column=col).alignment = Alignment(horizontal="left" if col in (2, 3) else "center")
            if i % 2 == 1:
                ws.cell(row=r, column=col).fill = PatternFill(start_color=CINZA, end_color=CINZA, fill_type="solid")
    top15_end = top15_start + 14

    # ── Contagem nacional Usa vs Não Usa novo hospital (gráfico pizza) ──
    pie_row = top15_end + 3
    ws.cell(row=pie_row, column=8, value="Usa o Novo Hospital").font = LABEL_FONT
    ws.cell(row=pie_row, column=9,
            value=f'=COUNTIF({S_IMPACTO}!I3:I{N_FREG_LAST_ROW},"Sim")')
    ws.cell(row=pie_row + 1, column=8, value="Mantém-se no H. de Faro").font = LABEL_FONT
    ws.cell(row=pie_row + 1, column=9,
            value=f'=COUNTIF({S_IMPACTO}!I3:I{N_FREG_LAST_ROW},"Não")')

    # ── Tabela: Capacidade vs Procura por hospital (gráfico de capacidade) ──
    cap_row = pie_row + 4
    ws.cell(row=cap_row, column=8, value="Hospital").font = LABEL_FONT
    ws.cell(row=cap_row, column=9, value="Capacidade (camas)").font = LABEL_FONT
    ws.cell(row=cap_row, column=10, value="Procura Atribuída (camas)").font = LABEL_FONT
    ws.cell(row=cap_row + 1, column=8, value="Hospital de Faro (existente)")
    ws.cell(row=cap_row + 1, column=9, value="='1. Resumo'!C24")
    ws.cell(row=cap_row + 1, column=10, value="='1. Resumo'!C27")
    ws.cell(row=cap_row + 2, column=8, value="Novo Hospital (Portimão)")
    ws.cell(row=cap_row + 2, column=9, value="='1. Resumo'!C25")
    ws.cell(row=cap_row + 2, column=10, value="='1. Resumo'!C26")

    # ══════════════════════════════════════════════════════════════════
    # GRÁFICOS — pilha vertical, largura inteira, sem sobreposições
    # ══════════════════════════════════════════════════════════════════
    CHART_W_WIDE = 24
    CHART_H = 11
    ROWS_PER_CHART = 24

    def legend_bottom(chart):
        chart.legend.position = "b"
        chart.legend.overlay = False

    next_row = max(conc_end_row, top10_end) + 3

    # 1) Barras horizontais: Tempo médio Atual vs C/Novo Hospital por Concelho
    style_section_title(ws, f"B{next_row}:J{next_row}",
                         "📈  GRÁFICO — Tempo médio de acesso por Concelho (Atual vs. C/ Novo Hospital)")
    chart1_row = next_row + 1
    chart1 = BarChart()
    chart1.type = "bar"
    chart1.grouping = "clustered"
    chart1.title = "Tempo médio ponderado por Concelho (min)"
    chart1.x_axis.title = "Minutos"
    chart1.y_axis.title = "Concelho"
    chart1.style = 10
    data1 = Reference(ws, min_col=4, max_col=5, min_row=hdr_row, max_row=conc_end_row)
    cats1 = Reference(ws, min_col=2, min_row=conc_start_row, max_row=conc_end_row)
    chart1.add_data(data1, titles_from_data=True)
    chart1.set_categories(cats1)
    chart1.width = CHART_W_WIDE
    chart1.height = CHART_H + 4
    legend_bottom(chart1)
    ws.add_chart(chart1, f"B{chart1_row}")
    next_row = chart1_row + ROWS_PER_CHART + 6

    # 2) Barras: Top 10 candidatos por tempo médio ponderado
    style_section_title(ws, f"B{next_row}:J{next_row}",
                         "📈  GRÁFICO — Top 10 candidatos a nova localização (tempo médio, min)")
    chart2_row = next_row + 1
    chart2 = BarChart()
    chart2.type = "bar"
    chart2.title = "Top 10 candidatos — tempo médio ponderado (min)"
    chart2.y_axis.title = "Freguesia candidata"
    chart2.x_axis.title = "Minutos"
    chart2.style = 11
    data2 = Reference(ws, min_col=9, max_col=9, min_row=hdr_row, max_row=top10_end)
    cats2 = Reference(ws, min_col=8, min_row=top10_start, max_row=top10_end)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    chart2.width = CHART_W_WIDE
    chart2.height = CHART_H
    legend_bottom(chart2)
    ws.add_chart(chart2, f"B{chart2_row}")
    next_row = chart2_row + ROWS_PER_CHART

    # 3) Barras horizontais: Top 15 freguesias com maior poupança
    style_section_title(ws, f"B{next_row}:J{next_row}",
                         "📈  GRÁFICO — Top 15 freguesias com maior poupança de tempo (min)")
    chart3_row = next_row + 1
    chart3 = BarChart()
    chart3.type = "bar"
    chart3.title = "Top 15 freguesias — poupança de tempo (min)"
    chart3.x_axis.title = "Minutos poupados"
    chart3.style = 12
    data3 = Reference(ws, min_col=6, max_col=6, min_row=top15_hdr, max_row=top15_end)
    cats3 = Reference(ws, min_col=2, min_row=top15_start, max_row=top15_end)
    chart3.add_data(data3, titles_from_data=True)
    chart3.set_categories(cats3)
    chart3.width = CHART_W_WIDE
    chart3.height = CHART_H + 4
    legend_bottom(chart3)
    ws.add_chart(chart3, f"B{chart3_row}")
    next_row = chart3_row + ROWS_PER_CHART + 6

    # 4) Pizza: Usa Novo Hospital vs Mantém H. de Faro
    style_section_title(ws, f"B{next_row}:E{next_row}",
                         "📈  GRÁFICO — Freguesias que passam a usar o novo hospital")
    chart4_row = next_row + 1
    chart4 = PieChart()
    chart4.title = "Novo Hospital vs. Hospital de Faro"
    data4 = Reference(ws, min_col=9, min_row=pie_row, max_row=pie_row + 1)
    cats4 = Reference(ws, min_col=8, min_row=pie_row, max_row=pie_row + 1)
    chart4.add_data(data4, titles_from_data=False)
    chart4.set_categories(cats4)
    chart4.dataLabels = DataLabelList()
    chart4.dataLabels.showPercent = True
    chart4.dataLabels.showVal = True
    chart4.width = 13
    chart4.height = CHART_H
    legend_bottom(chart4)
    ws.add_chart(chart4, f"B{chart4_row}")

    # 5) Barras: Capacidade vs Procura por hospital
    style_section_title(ws, f"G{next_row}:J{next_row}",
                         "📈  GRÁFICO — Capacidade vs. Procura atribuída (camas)")
    chart5 = BarChart()
    chart5.type = "col"
    chart5.grouping = "clustered"
    chart5.title = "Capacidade vs. Procura por hospital"
    chart5.y_axis.title = "Camas"
    chart5.style = 13
    data5 = Reference(ws, min_col=9, max_col=10, min_row=cap_row, max_row=cap_row + 2)
    cats5 = Reference(ws, min_col=8, min_row=cap_row + 1, max_row=cap_row + 2)
    chart5.add_data(data5, titles_from_data=True)
    chart5.set_categories(cats5)
    chart5.width = 13
    chart5.height = CHART_H
    legend_bottom(chart5)
    ws.add_chart(chart5, f"G{chart4_row}")
    next_row = chart4_row + ROWS_PER_CHART

    # esconder coluna auxiliar (N = lista concelhos)
    ws.column_dimensions["N"].hidden = True

    ws.freeze_panes = "B8"
    ws.sheet_view.zoomScale = 90
    wb.active = 0

    print("A guardar workbook...")
    wb.save(OUTPUT_FILE)
    print(f"✅ Guardado em: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()